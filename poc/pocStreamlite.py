import streamlit as st
import pandas as pd
from docx import Document
import os

# Bepaal het pad naar de huidige map
current_dir = os.path.dirname(__file__)

# Pad naar de bestanden
excel_path = os.path.join(current_dir, "testadressen.xlsx")
template_path = os.path.join(current_dir, "BriefSjabloon.docx")

# Laad de bestanden
try:
    df = pd.read_excel(excel_path)
    st.success("Excel-bestand geladen!")
except FileNotFoundError:
    st.error("Excel-bestand niet gevonden. Upload handmatig.")
    uploaded_excel = st.file_uploader("Upload Excel-bestand", type=["xlsx"])
    if uploaded_excel:
        df = pd.read_excel(uploaded_excel)

try:
    doc = Document(template_path)
    st.success("Word-sjabloon geladen!")
except FileNotFoundError:
    st.error("Word-sjabloon niet gevonden. Upload handmatig.")
    uploaded_template = st.file_uploader("Upload Word-sjabloon", type=["docx"])
    if uploaded_template:
        doc = Document(uploaded_template)

# Rest van je code (dropdown, document genereren, etc.)
afdelingen = df["Afdeling"].tolist()
geselecteerde_afdeling = st.selectbox("Kies een afdeling", afdelingen)

if geselecteerde_afdeling:
    afdeling_data = df[df["Afdeling"] == geselecteerde_afdeling].iloc[0].to_dict()

    # Vul het sjabloon
    for paragraph in doc.paragraphs:
        for key, value in afdeling_data.items():
            if f"[{key}]" in paragraph.text:
                paragraph.text = paragraph.text.replace(f"[{key}]", str(value))

    # Sla het gevulde document op
    output_path = os.path.join(current_dir, "gevuld_document.docx")
    doc.save(output_path)

    # Downloadknop
    with open(output_path, "rb") as f:
        st.download_button(
            label="Download gevuld document",
            data=f,
            file_name="gevuld_document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
